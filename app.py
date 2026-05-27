import streamlit as st
import openpyxl
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
import io, json, re, uuid, os
from PIL import Image as PILImage
from supabase import create_client, Client 
import streamlit as st
import time

# --- 1. Supabase 설정 ---
@st.cache_resource
def get_supabase() -> Client:
    url = st.secrets["supabase"]["url"]
    key = st.secrets["supabase"]["key"]
    return create_client(url, key)

supabase = get_supabase()
BUCKET_NAME = "templates"

# ---  Supabase DB 연동  ---
def save_log_to_db(action, filename, p_time=0):
    """상세 정보를 포함하여 로그를 저장합니다."""
    try:
        # 1. 접속 환경 정보 가져오기
        user_agent = st.context.headers.get("User-Agent", "Unknown Agent")
        
        # 2. IP 주소 가져오기 (배포 환경에 따라 다름)
        ip = st.context.headers.get("X-Forwarded-For")
        if not ip:
            ip = st.context.headers.get("Remote-Addr", "Unknown IP")
        else:
            ip = ip.split(",")[0] # 실제 사용자 IP만 추출

        # 3. 데이터 구성 (컬럼명이 DB와 정확히 일치해야 함)
        log_data = {
            "action_type": action,
            "target_filename": filename,
            "processing_time": round(p_time, 2), # 소수점 2자리 초
            "user_agent": user_agent,
            "user_ip": ip
        }
        
        # 4. Supabase 전송
        supabase.table("user_logs").insert(log_data).execute()
        
    except Exception as e:
        # 오류 발생 시 사용자 화면이 아닌 콘솔에만 기록 (사용자 방해 방지)
        print(f"Log Error: {e}")

def load_presets_from_db():
    """DB 테이블 'excel_presets'에서 프리셋 로드"""
    try:
        res = supabase.table("excel_presets").select("*").execute()
        presets = {}
        for row in res.data:
            presets[row['preset_name']] = {
                "filename": row['filename'],
                "cells": row['cell_positions'],
                "template_name": row['template_path']
            }
        return presets
    except Exception as e:
        return {}

def save_preset_to_db(name, filename, cells, t_name):
    """DB 테이블 'excel_presets'에 프리셋 저장 (Upsert)"""
    try:
        data = {
            "preset_name": name,
            "filename": filename,
            "cell_positions": cells,
            "template_path": t_name
        }
        supabase.table("excel_presets").upsert(data, on_conflict="preset_name").execute()
        return True
    except Exception as e:
        st.error(f"DB 저장 실패: {e}")
        return False

def delete_preset_from_db(preset_name):
    """DB 테이블에서 특정 프리셋 삭제"""
    try:
        supabase.table("excel_presets").delete().eq("preset_name", preset_name).execute()
        return True
    except Exception as e:
        st.error(f"삭제 실패: {e}")
        return False

def upload_template_to_supabase(file_name, file_data):
    try:
        extension = os.path.splitext(file_name)[1]
        safe_db_name = f"template_{uuid.uuid4().hex[:8]}{extension}" 
        supabase.storage.from_(BUCKET_NAME).upload(
            path=safe_db_name, 
            file=file_data,
            file_options={"cache-control": "3600", "upsert": "true"}
        )
        return safe_db_name
    except Exception as e:
        st.error(f"업로드 실패: {e}")
        return None

def download_template_from_supabase(file_name):
    try:
        return supabase.storage.from_(BUCKET_NAME).download(file_name)
    except Exception as e:
        return None

# --- 3. 이미지 처리 로직 ---
def fit_image_to_merged_cell(ws, img_data, cell_addr):
    try:
        try:
            from pillow_heif import register_heif_opener
            register_heif_opener()
        except ImportError: pass

        input_img = PILImage.open(io.BytesIO(img_data))
        if input_img.width > 1600:
            ratio = 1600 / float(input_img.width)
            hsize = int(float(input_img.height) * ratio)
            input_img = input_img.resize((1600, hsize), PILImage.Resampling.LANCZOS)
        
        img_byte_arr = io.BytesIO()
        input_img.convert("RGB").save(img_byte_arr, format='JPEG', quality=85)
        img_byte_arr.seek(0)
        
        img = Image(img_byte_arr)
        target_range = None
        for merged_range in ws.merged_cells.ranges:
            if cell_addr.upper() in merged_range:
                target_range = merged_range
                break
        
        if target_range:
            sc, sr = target_range.min_col - 1, target_range.min_row - 1
            ec, er = target_range.max_col, target_range.max_row
            img.anchor = TwoCellAnchor('twoCell', AnchorMarker(sc, 0, sr, 0), AnchorMarker(ec, 0, er, 0))
        else:
            img.anchor = cell_addr.upper()
        
        ws.add_image(img)
    except Exception as e:
        st.error(f"❌ 이미지 삽입 실패 ({cell_addr}): {e}")

# --- 4. 메인 UI ---
st.set_page_config(page_title="EXEL UPLOAD (by.Simroot)", layout="wide")
st.title("EXEL UPLOAD")

if 'presets' not in st.session_state:
    st.session_state.presets = load_presets_from_db()
if 'temp_cells' not in st.session_state:
    st.session_state.temp_cells = ""

# 프리셋 관리 사이드바
st.sidebar.header("프리셋 설정")
preset_list = list(st.session_state.presets.keys())
selected_preset = st.sidebar.selectbox("설정 불러오기", ["직접 입력"] + preset_list)

# 프리셋 데이터 매핑 로직
d_name, d_cells, active_temp_data, active_temp_name = "", "", None, None
if selected_preset != "직접 입력" and selected_preset in st.session_state.presets:
    p = st.session_state.presets[selected_preset]
    d_name = p.get("filename", "")
    d_cells = p.get("cells", "")
    st.session_state.temp_cells = d_cells
    t_name = p.get("template_name")
    if t_name:
        active_temp_data = download_template_from_supabase(t_name)
        if active_temp_data:
            active_temp_name = t_name
            st.sidebar.success(f"✅ 양식 연결됨: {t_name}")

if selected_preset != "직접 입력":
    if st.sidebar.button("현재 프리셋 삭제", width="stretch"):
        if delete_preset_from_db(selected_preset):
            st.sidebar.success(f"'{selected_preset}' 삭제 완료!")
            st.session_state.presets = load_presets_from_db()
            st.rerun()

# UI 레이아웃
main_col1, main_col2 = st.columns([1, 1])

with main_col1:
    st.subheader("설정")
    custom_filename = st.text_input("결과 파일명", value=d_name)
    
    st.write("**위치 설정**")
    w_col1, w_col2, w_col3 = st.columns([1, 1, 1.5])
    with w_col1: 
        s_num = st.number_input("시트", min_value=1, step=1)
    with w_col2: 
        all_cols = [chr(i) for i in range(ord('A'), ord('Z') + 1)]
        default_idx = all_cols.index("B") if "B" in all_cols else 0
        c_let = st.selectbox("열", all_cols, index=default_idx)
    with w_col3: 
        r_num = st.number_input("행 번호", min_value=1, step=1)
    
    cur_addr = f"{s_num}:{c_let}{r_num}"
    
    b_col1, b_col2, b_col3 = st.columns(3)
    with b_col1:
        if st.button("위치 추가", width="stretch"):
            if st.session_state.temp_cells: 
                st.session_state.temp_cells += f", {cur_addr}"
            else: 
                st.session_state.temp_cells = cur_addr
    with b_col2:
        if st.button("+16행 추가", width="stretch"): 
            st.session_state.temp_cells += "+16"
    with b_col3:
        if st.button("전체 초기화", width="stretch"):
            st.session_state.temp_cells = ""
            st.rerun()

    cell_input = st.text_area("최종 위치", value=st.session_state.temp_cells, height=100)
    st.session_state.temp_cells = cell_input
    uploaded_excel = st.file_uploader("엑셀 양식 업로드", type=['xlsx'])

    new_p_name = st.sidebar.text_input("신규 프리셋 이름")
    if st.sidebar.button("프리셋 저장", use_container_width=True):
        if new_p_name and custom_filename and cell_input:
            target_t_name = active_temp_name
            if uploaded_excel:
                excel_data = uploaded_excel.getvalue()
                safe_name = upload_template_to_supabase(uploaded_excel.name, excel_data)
                if safe_name: target_t_name = safe_name
            
            if save_preset_to_db(new_p_name, custom_filename, cell_input, target_t_name):
                st.sidebar.success(f"'{new_p_name}' 저장 완료!")
                st.session_state.presets = load_presets_from_db()
                st.rerun()

with main_col2:
    st.subheader("사진 업로드")
    uploaded_imgs = st.file_uploader("작업 사진 선택", type=['jpg', 'jpeg', 'png'], accept_multiple_files=True)
    
    if st.button("보고서 생성"):
        start_time = time.time()
        final_excel_data = uploaded_excel.getvalue() if uploaded_excel else active_temp_data
        
        final_cells = []
        parts = [p.strip() for p in cell_input.split(",") if p.strip()]
        for p in parts:
            if "+" in p:
                try:
                    base, inc = p.split("+")
                    s_idx, s_cell = base.split(":")
                    row_inc = int(inc)
                    match = re.match(r"([a-zA-Z]+)([0-9]+)", s_cell)
                    c_l, r_s = match.group(1), int(match.group(2))
                    for k in range(50): final_cells.append(f"{s_idx}:{c_l}{r_s + (k * row_inc)}")
                except: continue
            else:
                if ":" in p: final_cells.append(p)

        if final_excel_data and final_cells and uploaded_imgs:
            with st.spinner("엑셀 보고서를 생성 중입니다..."):
                try:
                    wb = openpyxl.load_workbook(io.BytesIO(final_excel_data))
                    img_bytes_list = [img_file.read() for img_file in uploaded_imgs]
                    for i, img_bytes in enumerate(img_bytes_list):
                        if i >= len(final_cells): break
                        if not img_bytes: continue
                        p_idx_str, cell_addr = final_cells[i].split(":")
                        ws = wb[wb.sheetnames[int(p_idx_str) - 1]]
                        end_time = time.time()
                        fit_image_to_merged_cell(ws, img_bytes, cell_addr)
                    
                    out = io.BytesIO()
                    wb.save(out)
                    out.seek(0)

                    end_time = time.time() #종료시간
                    duration = end_time - start_time
                    st.success("✅ 생성 완료!")
                    save_log_to_db("보고서 생성", f"{custom_filename}.xlsx", p_time=end_time - start_time)
                    st.download_button("결과 엑셀 다운로드", data=out, file_name=f"{custom_filename}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)
                except Exception as e: st.error(f"오류 발생: {e}")


# --- 관리자 전용  ---
for _ in range(15): 
    st.sidebar.write("")
st.sidebar.markdown("---")
admin_key = st.sidebar.checkbox("관리자 모드 접속")

if admin_key:
    password = st.sidebar.text_input("Admin Password", type="password")
    if password == "@tlavmf123":
        st.header(" Administrator DB ")
        
        # 탭을 사용하여 프리셋 관리와 이용 기록 분리
        tab_preset, tab_log = st.tabs(["프리셋 관리", "이용 기록 조회"])
        
        # --- 프리셋 관리 ---
        with tab_preset:
            try:
                res = supabase.table("excel_presets").select("*").execute()
                if res.data:
                    st.subheader("DB 저장 프리셋 목록")
                    st.dataframe(res.data, width="stretch")
                    st.divider()
                    col_del1, col_del2 = st.columns([3, 1])
                    with col_del1:
                        target_id = st.selectbox("삭제할 프리셋 ID 선택", [r['id'] for r in res.data], key="del_preset_sb")
                    with col_del2:
                        if st.button("행 삭제", type="secondary", width="stretch"):
                            supabase.table("excel_presets").delete().eq("id", target_id).execute()
                            st.success("삭제 완료!")
                            st.rerun()
                else:
                    st.info("저장된 프리셋이 없습니다.")
            except Exception as e:
                st.error(f"프리셋 로드 실패: {e}")

        # --- 이용 기록 조회 ---
        with tab_log:
            try:
                # 최신순 정렬 (desc=True)
                log_res = supabase.table("user_logs").select("*").order("created_at", desc=True).limit(50).execute()
                
                if log_res.data:
                    import pandas as pd
                    from datetime import timedelta 
                    df = pd.DataFrame(log_res.data)
                    
                    if 'created_at' in df.columns:
                        # 1. 먼저 datetime 형식으로 변환 (UTC 기준)
                        df['created_at'] = pd.to_datetime(df['created_at'])
                        
                        # 2. 한국 시간(UTC+9)으로 9시간
                        df['created_at'] = df['created_at'] + timedelta(hours=9)
                        df['created_at'] = df['created_at'].dt.strftime('%Y-%m-%d %H:%M:%S')
                    
                    st.subheader("최근 상세 이용 기록")
                    st.dataframe(df, width="stretch")
                    
                    st.divider()
                    if st.button("로그 기록 전체 삭제", type="primary", width="stretch"):
                        # 모든 로그 삭제
                        supabase.table("user_logs").delete().neq("id", 0).execute()
                        st.success("로그 초기화 완료")
                        st.rerun()
                else:
                    st.info("아직 생성된 이용 기록이 없습니다.")
            except Exception as e:
                st.error(f"로그 로드 실패: {e}")
                st.info("💡 SQL Editor에서 user_logs 테이블과 컬럼(processing_time, user_ip 등)이 생성되었는지 확인하세요.")

    elif password:
        st.sidebar.warning("비밀번호가 틀렸습니다.")

st.markdown("---")
st.info("""💡 **Supabase 연동으로 데이터가 안전하게 보존됩니다.**
- **자동증가:** `1:B2+16`으로 입력하면 B2, B18, B34 순으로 사진이 들어갑니다.
- **프리셋:** 한 번 저장해두면 다음부터는 사진만 올리면 끝납니다.
""")